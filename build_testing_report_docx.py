from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = r"D:\medclin\Raport_Capitolul_3_Testarea_MediClin.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, widths)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "E8EEF5")
        set_cell_text(cell, header, bold=True, color="1F3A5F")
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            set_cell_text(cells[idx], value)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.1
    doc.add_paragraph()
    return table


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "F4F6F9")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor(31, 58, 95)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Title", 22, "0B2545", 0, 8),
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)


def main():
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Capitolul 3. Implementarea si testarea sistemului MediClin").bold = True
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Material pentru raport: capitolele 3.1 si 3.2").italic = True
    subtitle.runs[0].font.color.rgb = RGBColor(85, 85, 85)

    add_callout(
        doc,
        "Scopul documentului",
        "Acest document poate fi inclus in raportul proiectului pentru a demonstra functionalitatea sistemului MediClin si modul in care aplicatia a fost testata manual si automat.",
    )

    doc.add_heading("3.1 Secvente de cod si rezultate obtinute din aplicatie", level=1)
    doc.add_paragraph(
        "In cadrul aplicatiei MediClin au fost implementate functionalitati pentru trei roluri principale: pacient, medic si administrator. "
        "Sistemul include autentificare, navigare pe roluri, fisa medicala electronica, retete, rezultate ale analizelor, programari, mesagerie, rapoarte si setari."
    )

    doc.add_heading("3.1.1 Exemple de module implementate", level=2)
    add_table(
        doc,
        ["Modul", "Secventa / componenta", "Rezultat obtinut"],
        [
            ["Autentificare", "LoginWindow + LoginViewModel", "Utilizatorul introduce email si parola, iar sistemul deschide interfata corespunzatoare rolului."],
            ["Pacient", "PatientMainWindow, MyEMRView, PrescriptionsView", "Pacientul poate consulta fisa medicala, retetele active si rezultatele analizelor."],
            ["Medic", "DoctorMainWindow, EMRView, PatientsView", "Medicul poate cauta pacienti, completa fisa medicala, adauga alergii, emite retete si transmite analize."],
            ["Administrator", "AdminMainWindow, UsersManagementView, GlobalSettingsView", "Administratorul poate gestiona utilizatori, medici, setari globale si informatii financiare."],
            ["Persistenta datelor", "Repository-uri pentru pacienti, consultatii, retete, analize si notificari", "Datele sunt salvate si citite din baza de date MySQL."],
        ],
        [1.35, 2.25, 2.9],
    )

    doc.add_heading("3.1.2 Rezultate obtinute in aplicatie", level=2)
    add_bullet(doc, "Dupa autentificare, utilizatorul este redirectionat catre interfata rolului sau.")
    add_bullet(doc, "Pacientul poate vizualiza istoricul consultatiilor si poate selecta o consultatie pentru completarea automata a campurilor din fisa.")
    add_bullet(doc, "Medicul poate adauga rezultate pentru analize, alergii, recomandari, diagnostic si retete.")
    add_bullet(doc, "Pacientul poate solicita reinnoirea unei retete, iar medicul de familie primeste notificare.")
    add_bullet(doc, "Administratorul poate modifica tema aplicatiei si culoarea primara.")

    doc.add_paragraph("Capturi recomandate pentru raport: login, dashboard pacient, fisa medicala pacient, lista pacienti medic, fisa medic, pagina financiara admin/medic, setari globale.")

    doc.add_heading("3.2 Testarea sistemului", level=1)
    doc.add_heading("3.2.1 Notiuni generale despre testare", level=2)
    doc.add_paragraph(
        "Testarea software reprezinta procesul prin care se verifica daca un sistem informatic functioneaza conform cerintelor stabilite. "
        "Prin testare se identifica erori, se valideaza functionalitatile si se confirma faptul ca aplicatia este stabila, usor de utilizat si pregatita pentru utilizare reala."
    )
    doc.add_paragraph(
        "In cazul sistemului MediClin, testarea este importanta deoarece aplicatia gestioneaza date sensibile: informatii despre pacienti, fise medicale, retete, rezultate ale analizelor si comunicarea dintre medic si pacient."
    )

    doc.add_heading("3.2.2 Importanta testarii", level=2)
    add_bullet(doc, "Confirma functionarea corecta a butoanelor, meniurilor si formularelor.")
    add_bullet(doc, "Verifica salvarea si afisarea corecta a datelor medicale.")
    add_bullet(doc, "Reduce riscul aparitiei erorilor in utilizarea aplicatiei.")
    add_bullet(doc, "Asigura o experienta mai buna pentru pacient, medic si administrator.")
    add_bullet(doc, "Ofera dovezi pentru raport prin rezultate, capturi de ecran si concluzii.")

    doc.add_heading("3.2.3 Modalitati de testare", level=2)
    add_table(
        doc,
        ["Modalitate", "Descriere", "Avantaje", "Limitari"],
        [
            ["Testare manuala", "Testerul verifica fiecare functionalitate prin actiuni directe in aplicatie.", "Usor de realizat, potrivita pentru interfata grafica si verificari vizuale.", "Consuma timp si poate fi repetitiva."],
            ["Testare automata", "Actiunile sunt executate de un script sau de un instrument precum Katalon Studio.", "Rapida, repetabila, genereaza rezultate clare ale executiei.", "Necesita configurare initiala si mentenanta a testelor."],
        ],
        [1.25, 2.05, 1.65, 1.55],
    )

    doc.add_heading("3.2.4 Alegerea metodei de testare", level=2)
    doc.add_paragraph(
        "Pentru proiectul MediClin se recomanda utilizarea combinata a testarii manuale si automate. "
        "Testarea manuala este utila pentru verificarea initiala a interfetei, iar testarea automata cu Katalon Studio este potrivita pentru repetarea scenariilor principale si pentru obtinerea rapida a rezultatelor executiei."
    )
    add_callout(
        doc,
        "Argument pentru Katalon Studio",
        "Katalon Studio poate fi folosit pentru testarea aplicatiilor Windows desktop, inclusiv aplicatii WPF. Instrumentul poate inregistra pasii testerului si ii poate rula ulterior automat, generand rapoarte cu rezultatele testarii.",
    )

    doc.add_heading("3.2.5 Testare manuala", level=2)
    doc.add_paragraph(
        "Pentru testarea manuala au fost accesate principalele butoane si pagini ale sistemului. Rezultatele sunt prezentate in tabelul urmator."
    )
    manual_rows = [
        ["Autentificare", "Email: pacient / medic / admin; parola valida", "Se apasa butonul Autentifica-te.", "Se deschide interfata rolului corect.", "Rolul corect se deschide.", "Acceptat"],
        ["Navigare pacient", "Cont pacient autentificat", "Se apasa Dashboard, Fisa medicala, Retete, Analize, Mesaje, Setari.", "Fiecare pagina se incarca fara eroare.", "Paginile se afiseaza corect.", "Acceptat"],
        ["Fisa medicala pacient", "Consultatii existente in istoric", "Se selecteaza o consultatie din istoric.", "Campurile fisei se completeaza cu datele consultatiei.", "Datele se afiseaza automat.", "Acceptat"],
        ["Retete pacient", "Reteta activa existenta", "Se apasa Solicita reinnoire pe reteta dorita.", "Medicul de familie primeste notificare.", "Solicitarea este trimisa.", "Acceptat"],
        ["Rezultate analize pacient", "Pacient autentificat", "Se acceseaza pagina Rezultate analize.", "Se afiseaza lista analizelor transmise.", "Pagina este disponibila.", "Acceptat"],
        ["Lista pacienti medic", "Cont medic autentificat", "Se acceseaza Pacienti si se cauta pacientul.", "Lista filtreaza pacientii si permite deschiderea fisei.", "Pacientul poate fi selectat.", "Acceptat"],
        ["Fisa medic", "Pacient selectat", "Medicul completeaza simptome, diagnostic, recomandari si salveaza.", "Datele sunt salvate in fisa medicala.", "Fisa este actualizata.", "Acceptat"],
        ["Adaugare alergii", "Pacient selectat", "Se introduce substanta, severitatea si se apasa Adauga.", "Alergia apare in lista pacientului.", "Alergia este adaugata.", "Acceptat"],
        ["Trimitere analiza", "Pacient selectat", "Se introduc laborator, test, valoare, status si se apasa Trimite analiza pacientului.", "Analiza apare la pacient si se creeaza notificare.", "Analiza este transmisa.", "Acceptat"],
        ["Financiar medic", "Cont medic autentificat", "Se acceseaza Financiar.", "Se afiseaza venitul estimat si activitatea medicului.", "Pagina este disponibila.", "Acceptat"],
        ["Setari admin", "Cont administrator", "Se modifica tema sau culoarea primara.", "Interfata isi schimba stilul conform setarii.", "Culoarea si tema se aplica.", "Acceptat"],
    ]
    add_table(
        doc,
        ["Functionalitate", "Date introduse", "Actiune", "Rezultat asteptat", "Rezultat primit", "Concluzie"],
        manual_rows,
        [1.15, 1.25, 1.2, 1.35, 1.15, 0.75],
    )

    doc.add_heading("3.2.6 Testare automata cu Katalon Studio", level=2)
    doc.add_paragraph(
        "Testarea automata se poate realiza in Katalon Studio prin inregistrarea pasilor efectuati manual. Dupa inregistrare, Katalon poate executa aceiasi pasi automat si poate genera rezultatul rularii."
    )
    add_table(
        doc,
        ["Nr.", "Scenariu automat", "Pasi inregistrati in Katalon", "Rezultat asteptat"],
        [
            ["1", "Autentificare pacient", "Pornire aplicatie, completare email/parola, apasare Autentifica-te.", "Se deschide meniul pacientului."],
            ["2", "Navigare in fisa medicala", "Apasare Fisa medicala, selectare consultatie din istoric.", "Campurile fisei se completeaza automat."],
            ["3", "Solicitare reinnoire reteta", "Deschidere Retete, apasare Solicita reinnoire pe un card.", "Apare mesaj de confirmare si se trimite notificare."],
            ["4", "Flux medic - pacient", "Login medic, deschidere Pacienti, deschidere fisa, adaugare analiza.", "Rezultatul analizei este transmis pacientului."],
            ["5", "Setari admin", "Login admin, deschidere Setari, schimbare culoare primara.", "Culoarea aplicatiei se schimba."],
        ],
        [0.45, 1.55, 2.75, 1.75],
    )

    doc.add_heading("3.2.7 Pasi pentru configurarea in Katalon Studio", level=2)
    steps = [
        "Se compileaza aplicatia MediClin si se identifica fisierul executabil Mediclin.UI.exe.",
        "In Katalon Studio se creeaza un proiect nou pentru testare Desktop/Windows.",
        "Se selecteaza aplicatia MediClin ca Application Under Test.",
        "Se porneste functia de Windows/Desktop Recorder.",
        "Se executa manual pasii doriti: login, navigare, completare formulare, salvare si verificare rezultat.",
        "Katalon salveaza pasii ca test case, iar testul poate fi rulat automat.",
        "La final se salveaza capturi de ecran cu rezultatele executiei pentru a fi introduse in raport.",
    ]
    for step in steps:
        p = doc.add_paragraph(style="List Number")
        p.add_run(step)

    doc.add_heading("3.2.8 Dovezi recomandate pentru raport", level=2)
    add_bullet(doc, "Captura cu test case-ul inregistrat in Katalon.")
    add_bullet(doc, "Captura cu executia testului: Passed / Failed.")
    add_bullet(doc, "Captura cu raportul generat de Katalon.")
    add_bullet(doc, "Capturi din aplicatie pentru rezultatul obtinut dupa test.")

    doc.add_heading("Concluzii", level=1)
    doc.add_paragraph(
        "In urma testarii, sistemul MediClin demonstreaza functionarea principalelor fluxuri pentru pacient, medic si administrator. "
        "Testarea manuala confirma functionalitatile vizibile ale interfetei, iar testarea automata cu Katalon Studio permite repetarea rapida a scenariilor principale si obtinerea unor rezultate clare pentru raport."
    )

    doc.add_heading("Bibliografie", level=1)
    add_bullet(doc, "Katalon Docs - Introduction to Desktop app testing in Katalon Studio.")
    add_bullet(doc, "Katalon Docs - Open a test project in Katalon Studio.")
    add_bullet(doc, "Documentatia proiectului MediClin si codul sursa al aplicatiei.")

    doc.core_properties.title = "Capitolul 3 - Testarea sistemului MediClin"
    doc.core_properties.author = "MediClin"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
