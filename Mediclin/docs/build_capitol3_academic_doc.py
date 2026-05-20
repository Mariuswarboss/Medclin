from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parent
OUT_PATH = OUT_DIR / "Capitolul_3_Testarea_sistemului_MediClin.docx"
FALLBACK_OUT_PATH = OUT_DIR / "Capitolul_3_Testarea_sistemului_MediClin_formatat.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="000000", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def style_document(document):
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(1)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1
    normal.paragraph_format.space_after = Pt(0)

    if "GeneralDiplomText" not in [s.name for s in document.styles]:
        general = document.styles.add_style("GeneralDiplomText", WD_STYLE_TYPE.PARAGRAPH)
    else:
        general = document.styles["GeneralDiplomText"]
    general.base_style = normal
    general.font.name = "Times New Roman"
    general._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    general.font.size = Pt(12)
    general.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    general.paragraph_format.first_line_indent = Cm(1.25)
    general.paragraph_format.line_spacing = 1.5
    general.paragraph_format.space_after = Pt(0)

    heading1 = document.styles["Heading 1"]
    heading1.font.name = "Times New Roman"
    heading1._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    heading1.font.size = Pt(13)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor(0, 0, 0)
    heading1.paragraph_format.space_before = Pt(30)
    heading1.paragraph_format.space_after = Pt(6)
    heading1.paragraph_format.line_spacing = 1.08

    heading2 = document.styles["Heading 2"]
    heading2.font.name = "Times New Roman"
    heading2._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    heading2.font.size = Pt(12)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor(0, 0, 0)
    heading2.paragraph_format.space_before = Pt(12)
    heading2.paragraph_format.space_after = Pt(6)

    if "CodSursa" not in [s.name for s in document.styles]:
        code = document.styles.add_style("CodSursa", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = document.styles["CodSursa"]
    code.base_style = normal
    code.font.name = "Courier New"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    code.font.size = Pt(9)
    code.paragraph_format.left_indent = Cm(0.5)
    code.paragraph_format.space_after = Pt(3)


def set_run_font(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def paragraph(document, text="", style="GeneralDiplomText", align=None):
    p = document.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, 12)
    return p


def heading(document, text, level=1, center=False):
    p = document.add_heading(text, level=level)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        set_run_font(run, 13 if level == 1 else 12, True)
    return p


def add_cover(document):
    for text in [
        "Ministerul Educației și Cercetării al Republicii Moldova",
        "Colegiul Universității Tehnice al Moldovei",
    ]:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, 14)

    for _ in range(5):
        document.add_paragraph()

    for text in ["RAPORTUL", "STAGIULUI DE PRACTICĂ"]:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, 20, True)

    document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("TEMA: Sistem de management pentru clinică")
    set_run_font(run, 18, True)

    for _ in range(3):
        document.add_paragraph()

    details = [
        "Elevul:\t\t  Chistrea Marius-Iustin",
        "Grupa:\t\t  PTPP-241",
        "Specialitatea:\t  Programarea și testarea produselor de program",
        "Baza de practică: Municipiul Chișinău",
    ]
    for text in details:
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Cm(4.5)
        run = p.add_run(text)
        set_run_font(run, 16)

    for _ in range(5):
        document.add_paragraph()

    for text, bold in [
        ("Conducătorul stagiului de practică de la", False),
        ("Colegiul Universității Tehnice a Moldovei", False),
        ("Pîntea Adina", True),
    ]:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(text)
        set_run_font(run, 14 if not bold else 16, bold)

    for _ in range(4):
        document.add_paragraph()

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Chișinău 2026")
    set_run_font(run, 14)
    document.add_page_break()


def add_table(document, title, headers, rows, widths=None):
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = caption.add_run(title)
    set_run_font(run, 12, True)

    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        set_cell_shading(cell, "EDEDED")
        set_cell_border(cell)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cells[idx].text = str(value)
            set_cell_border(cells[idx])
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = widths[idx]
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1
                for run in p.runs:
                    set_run_font(run, 10, row is table.rows[0])
    document.add_paragraph()
    return table


def add_code(document, title, lines):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(title)
    set_run_font(run, 12, True)
    for line in lines:
        code_p = document.add_paragraph(style="CodSursa")
        run = code_p.add_run(line)
        run.font.name = "Courier New"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        run.font.size = Pt(9)
    document.add_paragraph()


def build_document():
    document = Document()
    style_document(document)
    add_cover(document)

    p = document.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("REZUMAT")
    set_run_font(run, 13, True)
    paragraph(
        document,
        "Capitolul de față prezintă rezultatele implementării și testării aplicației MediClin, un sistem desktop destinat "
        "gestionării activității unei clinici medicale. Sunt descrise funcționalitățile principale realizate, rezultatele "
        "observate în aplicație și metodele de testare utilizate pentru validarea sistemului.",
    )
    paragraph(
        document,
        "Testarea a fost structurată în două direcții: testare manuală, prin parcurgerea fiecărui ecran și buton, și testare "
        "automată cu Katalon Studio, prin înregistrarea și reluarea scenariilor principale de lucru.",
    )
    document.add_page_break()

    heading(document, "CAPITOLUL 3. IMPLEMENTAREA ȘI TESTAREA SISTEMULUI MEDICLIN", center=True)
    paragraph(
        document,
        "În cadrul acestui capitol sunt prezentate secvențele principale de implementare, rezultatele obținute în aplicație "
        "și procesul de testare. Aplicația MediClin este organizată pe roluri: pacient, medic și administrator, fiecare rol "
        "având propriile pagini și funcționalități.",
    )

    heading(document, "3.1 Secvențe de cod și rezultate obținute din aplicație")
    paragraph(
        document,
        "Implementarea sistemului a urmărit realizarea unui flux complet de lucru pentru clinică: autentificare, navigare pe "
        "roluri, fișă medicală, istoric consultații, rețete, rezultate de analize, setări și zona financiară."
    )

    add_table(
        document,
        "Tabelul 3.1 - Module implementate în aplicația MediClin",
        ["Modul", "Rezultat obținut în aplicație"],
        [
            ("Autentificare și roluri", "Utilizatorul se conectează cu email și parolă, iar aplicația deschide zona corespunzătoare: pacient, medic sau administrator."),
            ("Fișă medicală", "Pacientul și medicul pot vizualiza simptome, diagnostic, semne vitale, recomandări și istoric consultații."),
            ("Istoric consultații", "La selectarea unei consultații din istoric, câmpurile fisei se completează automat cu informațiile consultației alese."),
            ("Rețete", "Medicul poate emite rețete, iar pacientul poate solicita reînnoirea unei rețete către medicul de familie."),
            ("Rezultate analize", "Medicul poate transmite rezultate de laborator, iar pacientul le vizualizează în pagina Rezultate analize."),
            ("Setări și design", "Aplicația permite schimbarea temei și alegerea culorii primare."),
            ("Financiar", "Administratorul și medicul pot consulta informații financiare estimative despre consultații și venituri."),
        ],
        [Cm(5), Cm(13)],
    )

    add_code(
        document,
        "Secvența 3.1 - Exemplu logică de completare automată a fișei",
        [
            "selectedConsultation = consultationHistory.SelectedItem;",
            "SymptomsTextBox.Text = selectedConsultation.Symptoms;",
            "DiagnosisTextBox.Text = selectedConsultation.Diagnosis;",
            "RecommendationsTextBox.Text = selectedConsultation.Recommendations;",
        ],
    )

    paragraph(
        document,
        "Rezultatul obținut este că pacientul poate selecta o consultație anterioară din istoric, iar formularul fișei medicale "
        "se completează fără introducere manuală repetată. Această funcționalitate reduce timpul de lucru și permite analizarea "
        "rapidă a datelor medicale existente."
    )

    add_table(
        document,
        "Tabelul 3.2 - Rezultate vizibile în aplicație",
        ["Funcționalitate", "Rezultat obținut", "Dovadă recomandată"],
        [
            ("Login", "După autentificare, utilizatorul ajunge în spațiul corespunzător contului său.", "Captură login + dashboard"),
            ("Pacient - Fișă medicală", "Câmpurile pentru simptome, diagnostic, semne vitale și recomandări sunt afișate organizat.", "Captură fișă medicală"),
            ("Pacient - Istoric", "O consultație selectată încarcă automat informațiile în formular.", "Captură istoric consultații"),
            ("Pacient - Rețete", "Rețetele active sunt listate și pot fi selectate pentru solicitare de reînnoire.", "Captură rețete active"),
            ("Medic - Pacienți", "Lista de pacienți permite deschiderea fișei și rămâne lizibilă în tema întunecată.", "Captură listă pacienți"),
            ("Medic - Analize", "Medicul completează și transmite rezultatele analizelor către pacient.", "Captură trimitere analize"),
            ("Admin - Financiar", "Sunt afișate consultațiile pe perioadă și estimările financiare.", "Captură financiar admin"),
        ],
        [Cm(4), Cm(9), Cm(5)],
    )

    heading(document, "3.2 Testarea sistemului")
    paragraph(
        document,
        "Testarea reprezintă procesul prin care se verifică dacă o aplicație funcționează conform cerințelor stabilite. "
        "Prin testare se identifică erori, se confirmă comportamentul corect al butoanelor și formularelor, se verifică "
        "validarea datelor și se crește încrederea că sistemul poate fi utilizat în condiții reale."
    )
    paragraph(
        document,
        "Testarea ajută la descoperirea defectelor înainte de prezentarea aplicației, confirmă că funcționalitățile răspund "
        "cerințelor utilizatorilor și oferă dovezi clare prin rezultate, rapoarte și capturi de ecran."
    )

    add_table(
        document,
        "Tabelul 3.3 - Modalități de testare",
        ["Modalitate", "Descriere", "Avantaje"],
        [
            ("Testare manuală", "Testerul parcurge fiecare ecran, apasă butoanele, introduce date și compară rezultatul primit cu cel așteptat.", "Este potrivită pentru verificări vizuale și pentru funcționalități noi."),
            ("Testare automată", "Un script sau o aplicație specializată repetă automat pașii definiți și raportează dacă testul a trecut sau a eșuat.", "Economisește timp, poate fi repetată rapid și generează rapoarte clare."),
        ],
        [Cm(4), Cm(9), Cm(5)],
    )

    heading(document, "Alegerea instrumentului Katalon Studio", level=2)
    paragraph(
        document,
        "Pentru testarea automată a fost aleasă aplicația Katalon Studio, deoarece permite înregistrarea pașilor realizați de "
        "utilizator, reluarea automată a testului și generarea unui raport cu rezultatele rulării. Katalon este potrivit pentru "
        "MediClin deoarece aplicația conține multe fluxuri care trebuie verificate repetat: autentificare, navigare, fișă "
        "medicală, rețete, analize și setări."
    )

    heading(document, "Testarea manuală", level=2)
    add_table(
        document,
        "Tabelul 3.4 - Testarea manuală a aplicației MediClin",
        ["Nr.", "Funcționalitate / buton", "Date introduse", "Rezultat așteptat", "Rezultat obținut", "Concluzie"],
        [
            ("1", "Login pacient", "Email și parolă valide", "Se deschide dashboard-ul pacientului.", "Dashboard pacient afișat.", "Admis"),
            ("2", "Login cu date greșite", "Email/parolă invalide", "Aplicația afișează mesaj de eroare.", "Autentificarea este refuzată.", "Admis"),
            ("3", "Navigare pacient", "Click pe toate paginile din meniu", "Fiecare pagină se deschide fără blocare.", "Paginile sunt accesibile.", "Admis"),
            ("4", "Selectare consultație din istoric", "Click pe o consultație existentă", "Câmpurile fișei se completează automat.", "Datele consultației apar în boxuri.", "Admis"),
            ("5", "Reînnoire rețetă", "Selectare rețetă și solicitare reînnoire", "Medicul de familie vede cererea.", "Cererea apare la medic.", "Admis"),
            ("6", "Login medic", "Cont medic valid", "Se deschide workspace-ul medicului.", "Workspace medic afișat.", "Admis"),
            ("7", "Căutare pacient", "Text introdus în câmpul de căutare", "Lista se filtrează după pacient.", "Pacientul căutat rămâne vizibil.", "Admis"),
            ("8", "Adăugare alergie", "Denumire alergie și salvare", "Alergia apare în lista pacientului.", "Alergia este memorată.", "Admis"),
            ("9", "Transmitere rezultate analize", "Valori laborator completate de medic", "Pacientul vede rezultatele la Analize.", "Rezultatele sunt afișate la pacient.", "Admis"),
            ("10", "Pagina financiară", "Accesare Financiar", "Se afișează estimările financiare.", "Datele financiare sunt vizibile.", "Admis"),
        ],
        [Cm(1), Cm(3.7), Cm(3.4), Cm(4.2), Cm(4), Cm(2)],
    )

    heading(document, "Testarea automată cu Katalon Studio", level=2)
    paragraph(
        document,
        "Pentru utilizarea Katalon Studio se creează un proiect nou, se alege testarea pentru aplicații Windows/Desktop, iar "
        "la configurarea aplicației se selectează executabilul Mediclin.UI.exe. Ulterior se pornește recorder-ul, se efectuează "
        "pașii doriti în aplicație, se salvează test case-ul și se rulează testul."
    )
    add_table(
        document,
        "Tabelul 3.5 - Scenarii automate recomandate pentru Katalon",
        ["Test automat", "Pași înregistrați", "Validare", "Rezultat raport Katalon"],
        [
            ("Autentificare pacient", "Deschide aplicația, completează email/parolă, apasă Autentifică-te.", "Verifică apariția numelui pacientului în dashboard.", "Passed/Failed"),
            ("Navigare pacient", "Apasă pe Fișă medicală, Rețete active, Rezultate analize și Setări.", "Verifică titlul fiecărei pagini.", "Passed/Failed"),
            ("Reînnoire rețetă", "Deschide Rețete active, selectează rețeta, apasă Solicita reînnoire.", "Verifică trimiterea cererii către medic.", "Passed/Failed"),
            ("Flux medic", "Login medic, deschide Pacienți, caută pacient, deschide fișa.", "Verifică formularul pacientului și istoricul.", "Passed/Failed"),
            ("Trimitere analiză", "Medicul completează rezultatul unei analize și îl trimite.", "Verifică apariția rezultatului în contul pacientului.", "Passed/Failed"),
        ],
        [Cm(4), Cm(6), Cm(5.5), Cm(2.5)],
    )

    add_table(
        document,
        "Tabelul 3.6 - Capturi necesare pentru raport",
        ["Nr.", "Captură necesară", "Ce trebuie să demonstreze"],
        [
            ("1", "Katalon Studio - test case înregistrat", "Se vede lista de pași automatizați."),
            ("2", "Rulare test suite", "Se vede execuția testului."),
            ("3", "Raport final", "Se vede statusul Passed/Failed și durata testului."),
            ("4", "Aplicația MediClin în timpul testului", "Se vede ecranul verificat automat."),
        ],
        [Cm(2), Cm(7), Cm(9)],
    )

    heading(document, "Concluzii", center=True)
    paragraph(
        document,
        "În urma testării, funcționalitățile principale ale sistemului MediClin pot fi validate prin două metode: manual, prin "
        "parcurgerea directă a fiecărui buton și formular, și automat, prin rularea scenariilor în Katalon Studio. Testarea "
        "manuală oferă control vizual asupra comportamentului aplicației, iar testarea automată permite repetarea rapidă a "
        "scenariilor și obținerea rapoartelor necesare pentru documentarea rezultatelor."
    )
    paragraph(
        document,
        "Prin combinarea celor două metode, raportul demonstrează că aplicația este verificată atât la nivel funcțional, cât și "
        "la nivel de experiență a utilizatorului. Capturile generate din Katalon Studio vor completa dovezile practice pentru "
        "capitolul 3.2."
    )

    section = document.sections[-1]
    add_page_number(section.footer.paragraphs[0])

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    try:
        document.save(OUT_PATH)
        return OUT_PATH
    except PermissionError:
        document.save(FALLBACK_OUT_PATH)
        return FALLBACK_OUT_PATH


if __name__ == "__main__":
    print(build_document())
