from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parent
OUT_PATH = OUT_DIR / "Capitolul_3_Testarea_sistemului_MediClin.docx"


PRIMARY = "0B7A61"
PRIMARY_DARK = "075846"
INK = RGBColor(15, 23, 42)
MUTED = RGBColor(100, 116, 139)
LINE = "D9E2EC"
SOFT = "F3F7FA"
ACCENT_SOFT = "E7F7F2"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LINE, size="8"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_style(table, header_fill=PRIMARY, header_font="FFFFFF"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = INK
            if i == 0:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor.from_string(header_font)
            elif i % 2 == 0:
                set_cell_shading(cell, "FAFCFE")
            else:
                set_cell_shading(cell, "FFFFFF")
    set_repeat_table_header(table.rows[0])


def add_table(document, headers, rows, widths=None, header_fill=PRIMARY):
    table = document.add_table(rows=1, cols=len(headers))
    hdr_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        hdr_cells[idx].text = header
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = width
    set_table_style(table, header_fill=header_fill)
    document.add_paragraph()
    return table


def add_note_box(document, title, body):
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, ACCENT_SOFT)
    set_cell_border(cell, "B6E7D8")
    paragraph = cell.paragraphs[0]
    title_run = paragraph.add_run(title)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor.from_string(PRIMARY_DARK)
    title_run.font.size = Pt(10.5)
    paragraph.add_run("\n")
    body_run = paragraph.add_run(body)
    body_run.font.color.rgb = INK
    body_run.font.size = Pt(10)
    document.add_paragraph()


def add_bullet(document, text):
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def add_numbered(document, text):
    p = document.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def style_document(document):
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color in (
        ("Title", 22, PRIMARY_DARK),
        ("Heading 1", 17, PRIMARY_DARK),
        ("Heading 2", 14, PRIMARY),
        ("Heading 3", 12, INK),
    ):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        if isinstance(color, str):
            style.font.color.rgb = RGBColor.from_string(color)
        else:
            style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)


def add_header_footer(document):
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.text = "MediClin | Capitolul 3 - Implementare si testare"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Raport de testare - versiune pentru prezentare")
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED


def build_document():
    document = Document()
    style_document(document)
    add_header_footer(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Capitolul 3. Implementarea si testarea sistemului MediClin")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Raport pentru aplicatia de management clinic")
    run.font.color.rgb = MUTED
    run.font.size = Pt(12)

    meta_rows = [
        ("Aplicatie", "MediClin"),
        ("Tip sistem", "Sistem integrat pentru pacient, medic si administrator"),
        ("Data raportului", "19.05.2026"),
        ("Obiectiv", "Prezentarea functionalitatilor realizate si a testarii manuale/automate"),
    ]
    add_table(document, ["Camp", "Descriere"], meta_rows, [Cm(4), Cm(12)], header_fill=PRIMARY_DARK)

    document.add_heading("3.1 Secvente de cod si rezultate obtinute din aplicatie", level=1)
    document.add_paragraph(
        "In cadrul aplicatiei MediClin au fost dezvoltate module separate pentru pacient, medic si administrator. "
        "Functionalitatile urmaresc fluxul real al unei clinici: autentificare, acces pe roluri, programari, fise medicale, "
        "retete, rezultate de analize, mesagerie si administrare financiara."
    )

    add_note_box(
        document,
        "Observatie pentru raport",
        "In aceasta sectiune se pot atasa capturi din aplicatie langa fiecare functionalitate: ecranul de login, "
        "dashboard-ul pacientului, fisa medicala, pagina medicului si consola de administrare.",
    )

    document.add_heading("Module implementate", level=2)
    modules = [
        ("Autentificare si roluri", "Utilizatorul se conecteaza cu email si parola, iar aplicatia deschide zona potrivita: pacient, medic sau administrator."),
        ("Fisa medicala", "Pacientul si medicul pot vizualiza informatii clinice, diagnostice, recomandari, semne vitale si istoricul consultarilor."),
        ("Istoric consultatii", "La selectarea unei consultatii din istoric, campurile fisei sunt completate automat cu datele consultatiei selectate."),
        ("Retete", "Medicul poate emite retete, iar pacientul poate solicita reinnoirea unei retete catre medicul de familie."),
        ("Rezultate analize", "Medicul poate transmite rezultate de laborator, iar pacientul le poate vedea in pagina de rezultate analize."),
        ("Setari si design", "Aplicatia permite tema deschisa/intunecata si alegerea culorii primare pentru un aspect vizual coerent."),
        ("Financiar", "Administratorul si medicul au acces la date financiare estimative pentru consultatii si venituri."),
    ]
    add_table(document, ["Modul", "Rezultat obtinut in aplicatie"], modules, [Cm(5), Cm(11)])

    document.add_heading("Exemple de rezultate vizibile", level=2)
    result_rows = [
        ("Login", "Dupa autentificare, utilizatorul ajunge in spatiul corespunzator contului sau.", "Captura login + dashboard"),
        ("Pacient - Fisa medicala", "Campurile pentru simptome, diagnostic, semne vitale si recomandari sunt afisate organizat.", "Captura fisa medicala"),
        ("Pacient - Istoric", "O consultatie selectata incarca automat informatiile in formular.", "Captura istoric consultatii"),
        ("Pacient - Retete", "Retetele active sunt listate si pot fi selectate pentru solicitare de reinnoire.", "Captura retete active"),
        ("Medic - Pacienti", "Lista de pacienti ramane lizibila in tema intunecata si permite deschiderea fisei.", "Captura lista pacienti"),
        ("Medic - Analize", "Medicul poate completa si trimite rezultatele analizelor catre pacient.", "Captura trimitere analize"),
        ("Admin - Financiar", "Sunt afisate consultatiile pe perioada si estimarile financiare.", "Captura financiar admin"),
    ]
    add_table(document, ["Functionalitate", "Rezultat obtinut", "Dovada recomandata"], result_rows, [Cm(4), Cm(8), Cm(4)])

    document.add_heading("3.2 Testarea sistemului", level=1)
    document.add_paragraph(
        "Testarea reprezinta procesul prin care se verifica daca o aplicatie functioneaza conform cerintelor stabilite. "
        "Prin testare se identifica erori, se confirma comportamentul corect al butoanelor si formularelor, se verifica "
        "validarea datelor si se creste increderea ca sistemul poate fi utilizat in conditii reale."
    )

    document.add_heading("Rolul testarii", level=2)
    for item in [
        "ajuta la descoperirea defectelor inainte ca aplicatia sa fie prezentata sau utilizata;",
        "confirma ca functionalitatile implementate raspund cerintelor utilizatorilor;",
        "reduce riscul de pierdere sau afisare gresita a datelor medicale;",
        "ofera dovezi clare prin rezultate, rapoarte si capturi de ecran;",
        "imbunatateste calitatea si stabilitatea sistemului.",
    ]:
        add_bullet(document, item)

    document.add_heading("Modalitati de testare", level=2)
    document.add_paragraph(
        "Exista mai multe modalitati de testare, insa pentru acest proiect sunt relevante doua metode principale: "
        "testarea manuala si testarea automata."
    )
    add_table(
        document,
        ["Modalitate", "Descriere", "Avantaje"],
        [
            ("Testare manuala", "Testerul parcurge fiecare ecran si apasa butoanele aplicatiei, introducand date si comparand rezultatul primit cu cel asteptat.", "Este usor de inteles, potrivita pentru verificari vizuale si pentru functionalitati noi."),
            ("Testare automata", "Un script sau o aplicatie specializata repeta automat pasii definiti si raporteaza daca testul a trecut sau a esuat.", "Economiseste timp, poate fi repetata rapid si genereaza rapoarte clare."),
        ],
        [Cm(3.5), Cm(8), Cm(4.5)],
    )

    document.add_heading("Argumentarea alegerii Katalon", level=2)
    document.add_paragraph(
        "Pentru testarea automata a fost aleasa aplicatia Katalon Studio, deoarece permite inregistrarea pasilor efectuati "
        "de utilizator, reluarea automata a testului si generarea unui raport cu rezultatele rularii. Aceasta metoda este "
        "potrivita pentru MediClin deoarece aplicatia are multe ecrane si butoane care trebuie verificate repetat: login, "
        "navigare, fisa medicala, retete, analize si setari."
    )
    add_note_box(
        document,
        "Capturi recomandate pentru raport",
        "Atasati o captura cu testul Katalon inregistrat, o captura cu rularea testului si o captura cu raportul final "
        "in care apare statusul Passed/Failed.",
    )

    document.add_heading("Testarea manuala", level=2)
    manual_rows = [
        ("1", "Login pacient", "Email si parola valide", "Se deschide dashboard-ul pacientului.", "Dashboard pacient afisat.", "Admis"),
        ("2", "Login cu date gresite", "Email/parola invalide", "Aplicatia afiseaza mesaj de eroare.", "Autentificarea este refuzata.", "Admis"),
        ("3", "Navigare pacient", "Click pe Dashboard, Programari, Fisa, Retete, Analize, Mesaje, Setari", "Fiecare pagina se deschide fara blocare.", "Paginile sunt accesibile.", "Admis"),
        ("4", "Selectare consultatie din istoric", "Click pe o consultatie existenta", "Campurile fisei se completeaza automat.", "Datele consultatiei apar in boxuri.", "Admis"),
        ("5", "Solicitare fisa medicala", "Click pe solicitare fisa de la medic", "Medicul primeste solicitarea.", "Solicitarea este trimisa.", "Admis"),
        ("6", "Reinnoire reteta", "Selectare reteta si click pe solicitare reinnoire", "Medicul de familie vede cererea.", "Cererea apare la medic.", "Admis"),
        ("7", "Login medic", "Cont medic valid", "Se deschide workspace-ul medicului.", "Workspace medic afisat.", "Admis"),
        ("8", "Cautare pacient", "Text introdus in campul de cautare", "Lista se filtreaza dupa pacient.", "Pacientul cautat ramane vizibil.", "Admis"),
        ("9", "Deschidere fisa pacient", "Click pe Deschide fisa", "Se afiseaza formularul clinic al pacientului.", "Fisa se deschide.", "Admis"),
        ("10", "Adaugare alergie", "Denumire alergie si salvare", "Alergia apare in lista pacientului.", "Alergia este memorata.", "Admis"),
        ("11", "Transmitere rezultate analize", "Valori laborator completate de medic", "Pacientul vede rezultatele la Analize.", "Rezultatele sunt afisate la pacient.", "Admis"),
        ("12", "Emiterea retetei", "Medicament, cantitate, durata", "Reteta apare in contul pacientului.", "Reteta este listata la pacient.", "Admis"),
        ("13", "Login administrator", "Cont admin valid", "Se deschide consola de administrare.", "Admin console afisat.", "Admis"),
        ("14", "Schimbare tema si culoare primara", "Selectare tema neagra/alba si culoare", "Interfata se actualizeaza coerent.", "Culoarea si tema sunt aplicate.", "Admis"),
        ("15", "Pagina financiara", "Accesare Financiar", "Se afiseaza estimarile financiare si consultatiile.", "Datele financiare sunt vizibile.", "Admis"),
    ]
    add_table(
        document,
        ["Nr.", "Functionalitate / buton", "Date introduse", "Rezultat asteptat", "Rezultat obtinut", "Concluzie"],
        manual_rows,
        [Cm(1), Cm(3.5), Cm(3.2), Cm(4), Cm(4), Cm(2)],
    )

    document.add_heading("Testarea automata cu Katalon Studio", level=2)
    document.add_paragraph(
        "Testarea automata se realizeaza prin inregistrarea actiunilor utilizatorului si rularea lor automata. "
        "Pentru aplicatia MediClin, scenariile importante sunt autentificarea, navigarea intre pagini si salvarea datelor."
    )

    document.add_heading("Pregatirea aplicatiei pentru Katalon", level=3)
    for step in [
        "Se construieste aplicatia MediClin in Visual Studio sau prin comanda de publicare.",
        "Se identifica fisierul executabil al aplicatiei, de exemplu Mediclin.UI.exe.",
        "Se deschide Katalon Studio si se creeaza un proiect nou de tip desktop/windows application.",
        "Se seteaza calea catre executabilul MediClin.",
        "Se porneste inregistrarea, se efectueaza pasii in aplicatie, apoi se salveaza test case-ul.",
        "Se ruleaza testul si se salveaza raportul cu rezultatele.",
    ]:
        add_numbered(document, step)

    document.add_heading("Scenarii automate recomandate", level=3)
    katalon_rows = [
        ("Autentificare pacient", "Deschide aplicatia, completeaza email/parola, apasa Autentificati-va.", "Verifica aparitia numelui pacientului in dashboard.", "Passed/Failed in raport"),
        ("Navigare pacient", "Apasa pe Fisa medicala, Retete active, Rezultate analize si Setari.", "Verifica titlul fiecarei pagini.", "Passed/Failed in raport"),
        ("Reinnoire reteta", "Deschide Retete active, selecteaza reteta, apasa Solicita reinnoire.", "Verifica mesajul/cererea transmisa medicului.", "Passed/Failed in raport"),
        ("Flux medic", "Login medic, deschide Pacienti, cauta pacient, deschide fisa.", "Verifica formularul pacientului si lista de istoric.", "Passed/Failed in raport"),
        ("Trimitere analiza", "Medic completeaza rezultatul unei analize si trimite catre pacient.", "Verifica aparitia rezultatului in contul pacientului.", "Passed/Failed in raport"),
        ("Setari tema", "Admin schimba tema si culoarea primara, apoi salveaza.", "Verifica schimbarea culorii si pastrarea setarii.", "Passed/Failed in raport"),
    ]
    add_table(
        document,
        ["Test automat", "Pasi inregistrati", "Validare", "Rezultat raport Katalon"],
        katalon_rows,
        [Cm(3.5), Cm(6), Cm(5), Cm(3)],
    )

    document.add_heading("Zone pentru capturi de ecran", level=2)
    capture_rows = [
        ("Captura 1", "Katalon Studio - test case inregistrat", "Se vede lista de pasi automatizati."),
        ("Captura 2", "Rulare test suite", "Se vede executia testului."),
        ("Captura 3", "Raport final", "Se vede statusul Passed/Failed si durata testului."),
        ("Captura 4", "Aplicatia MediClin in timpul testului", "Se vede ecranul verificat automat."),
    ]
    add_table(document, ["Nr.", "Captura necesara", "Ce trebuie sa demonstreze"], capture_rows, [Cm(2.5), Cm(6.5), Cm(7)])

    document.add_heading("Concluzii", level=1)
    document.add_paragraph(
        "In urma testarii, functionalitatile principale ale sistemului MediClin pot fi validate prin doua metode: manual, "
        "prin parcurgerea directa a fiecarui buton si formular, si automat, prin rularea scenariilor in Katalon Studio. "
        "Testarea manuala ofera control vizual asupra comportamentului aplicatiei, iar testarea automata permite repetarea "
        "rapida a scenariilor si obtinerea rapoartelor necesare pentru documentarea rezultatelor."
    )
    document.add_paragraph(
        "Prin combinarea celor doua metode, raportul demonstreaza ca aplicatia este verificata atat la nivel functional, "
        "cat si la nivel de experienta a utilizatorului. Capturile din Katalon vor completa dovezile practice pentru "
        "capitolul 3.2."
    )

    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_heading("Anexa A. Model scurt pentru descrierea unei testari manuale", level=1)
    document.add_paragraph(
        "Exemplu: Pentru butonul Autentificati-va au fost introduse datele unui cont valid de pacient. Rezultatul asteptat "
        "a fost deschiderea dashboard-ului pacientului. Rezultatul primit a fost afisarea paginii pacientului cu numele "
        "utilizatorului si meniul principal. Concluzia: test admis, functionalitatea lucreaza corect."
    )

    document.add_heading("Anexa B. Text scurt pentru includerea in raport dupa rularea Katalon", level=1)
    document.add_paragraph(
        "Scenariul automat a fost rulat in Katalon Studio. Aplicatia MediClin a fost deschisa automat, campurile au fost "
        "completate conform pasilor inregistrati, iar rezultatele au fost comparate cu validarile stabilite. In raportul "
        "Katalon, scenariul a primit statusul Passed, ceea ce confirma ca fluxul testat functioneaza conform asteptarilor."
    )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    document.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    print(build_document())
