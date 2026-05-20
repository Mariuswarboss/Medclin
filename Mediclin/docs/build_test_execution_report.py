from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parent
OUT_PATH = OUT_DIR / "Raport_Testare_Executata_MediClin.docx"


def set_font(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def paragraph(doc, text="", align=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.line_spacing = 1.5
    p.alignment = align or WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_font(run)
    return p


def heading(doc, text, level=1, center=False):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        set_font(run, 13 if level == 1 else 12, True)
    return p


def add_table(doc, title, headers, rows):
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = cap.add_run(title)
    set_font(r, 12, True)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header

    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)

    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_font(run, 10, row_index == 0)

    doc.add_paragraph()
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(1)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(12)

    heading(doc, "RAPORT DE TESTARE EXECUTATĂ - MEDICLIN", center=True)
    paragraph(
        doc,
        "Acest raport prezintă testarea efectuată pentru aplicația MediClin în data de 19.05.2026. "
        "Au fost verificate pornirea aplicației, conexiunea cu baza de date, testele automate existente și posibilitatea de rulare Katalon.",
    )

    heading(doc, "1. Mediul de testare")
    add_table(
        doc,
        "Tabelul 1 - Configurația mediului de testare",
        ["Element", "Valoare"],
        [
            ("Sistem testat", "MediClin - aplicație desktop Windows/WPF"),
            ("Executabil", r"D:\medclin\Mediclin\Mediclin.UI\bin\Debug\net8.0-windows\Mediclin.UI.exe"),
            ("Bază de date", "MySQL, baza mediclin_db"),
            ("Metode", "Testare manuală, testare automată locală, verificare pregătire Katalon"),
            ("Instrument automat disponibil", "dotnet vstest"),
            ("Katalon Studio", "Nu a fost găsit instalat în mediul local"),
        ],
    )

    heading(doc, "2. Testarea manuală executată")
    paragraph(
        doc,
        "Testarea manuală a fost pornită prin lansarea aplicației și verificarea ferestrelor afișate. "
        "Aplicația se lansează, însă fluxurile după autentificare sunt blocate de mesajul de conexiune la baza de date.",
    )
    add_table(
        doc,
        "Tabelul 2 - Rezultatele testării manuale",
        ["ID", "Acțiune testată", "Rezultat așteptat", "Rezultat obținut", "Status"],
        [
            ("M-01", "Pornirea aplicației MediClin", "Aplicația pornește și afișează interfața inițială.", "Procesul Mediclin.UI pornește și rămâne activ.", "Admis parțial"),
            ("M-02", "Verificarea ferestrei de conexiune", "Aplicația se conectează la baza de date și permite login.", "Apare fereastra «Eroare conexiune».", "Respins"),
            ("M-03", "Citirea mesajului de eroare", "În caz de eroare, mesajul trebuie să fie explicit.", "Mesaj afișat: Nu s-a putut conecta la baza de date mediclin_db. Detalii: Authentication failed. No credentials are available in the security package.", "Admis"),
            ("M-04", "Login pacient", "Se deschide dashboard-ul pacientului.", "Nu s-a putut testa, fiind blocat de eroarea de conexiune.", "Blocat"),
            ("M-05", "Navigare pacient", "Paginile pacientului se deschid din meniu.", "Nu s-a putut testa, deoarece login-ul nu este accesibil.", "Blocat"),
            ("M-06", "Flux medic", "Medicul poate deschide pacienți, fișe, analize și financiar.", "Nu s-a putut testa end-to-end, deoarece aplicația nu trece de conexiunea DB.", "Blocat"),
            ("M-07", "Flux administrator", "Administratorul poate accesa dashboard, utilizatori, financiar și setări.", "Nu s-a putut testa end-to-end, deoarece aplicația nu trece de conexiunea DB.", "Blocat"),
        ],
    )

    heading(doc, "3. Verificarea bazei de date")
    paragraph(
        doc,
        "Pentru a separa problema aplicației de problema serverului MySQL, conexiunea la server a fost verificată separat. "
        "Serverul MySQL rulează și portul 3306 răspunde, iar baza de date mediclin_db există.",
    )
    add_table(
        doc,
        "Tabelul 3 - Verificări bază de date",
        ["Verificare", "Rezultat"],
        [
            ("Serviciu MySQL", "MySQL80 rulează"),
            ("Port 3306", "TcpTestSucceeded = True"),
            ("Baza de date mediclin_db", "Există"),
            ("Tabel utilizatori", "7 utilizatori existenți"),
            ("Concluzie", "Serverul răspunde, dar aplicația WPF eșuează la autentificarea MySQL."),
        ],
    )

    heading(doc, "4. Testarea automată executată")
    paragraph(
        doc,
        "Testarea automată locală a fost executată direct pe assembly-ul de teste deja compilat, deoarece rebuild-ul complet al soluției "
        "este blocat în mediul curent de accesul interzis la directorul Microsoft SDKs.",
    )
    add_table(
        doc,
        "Tabelul 4 - Rezultate teste automate",
        ["Test", "Rezultat"],
        [
            ("LoginAsync_ReturnsUser_WhenCredentialsAreValid", "Passed"),
            ("LoginAsync_ReturnsNull_WhenParolaHashIsInvalid_DoesNotThrow", "Passed"),
            ("LoginAsync_ReturnsNull_WhenPasswordInvalid", "Passed"),
            ("RegisterAsync_ReturnsError_WhenEmailExists", "Passed"),
            ("Total", "4 Passed, 0 Failed, 0 Skipped"),
        ],
    )

    heading(doc, "5. Testarea automată cu Katalon")
    paragraph(
        doc,
        "Katalon Studio nu a fost găsit instalat în mediul local, de aceea nu a putut fi generat un raport real Katalon. "
        "Totuși, a fost pregătit scriptul Katalon și lista de cazuri manuale pentru import sau copiere în proiectul Katalon.",
    )
    add_table(
        doc,
        "Tabelul 5 - Fișiere pregătite pentru Katalon",
        ["Fișier", "Rol"],
        [
            (r"D:\medclin\Mediclin\tests\katalon\MediClin_Automated_TestCase.groovy", "Schelet test automat Katalon pentru aplicația Windows/WPF"),
            (r"D:\medclin\Mediclin\tests\katalon\MediClin_Manual_TestCases.csv", "Tabel cu cazuri de testare manuală"),
            (r"D:\medclin\Mediclin\docs\Script_Testare_MediClin.md", "Script complet cu pași manuali și automați"),
        ],
    )

    heading(doc, "6. Concluzie")
    paragraph(
        doc,
        "Testarea automată locală a trecut cu succes pentru cele patru teste existente. Testarea manuală end-to-end a aplicației "
        "a fost blocată la conexiunea MySQL din interfața WPF, deși serverul MySQL și baza de date sunt disponibile. "
        "Pentru finalizarea testării complete este necesară corectarea autentificării MySQL din aplicație, după care se pot rula scenariile Katalon pregătite.",
    )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    print(build())
